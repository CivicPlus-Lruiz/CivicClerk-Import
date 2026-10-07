"""
Meeting Bodies tab (Standard Import)
"""

import customtkinter as ctk
from tkinter import filedialog, messagebox
import os
from pathlib import Path
from typing import TYPE_CHECKING
import openpyxl
from gui.models import MeetingBody, Subtype
if TYPE_CHECKING:
    from gui.meeting_import_app import MeetingImportGUI

# ================================================================
# TAB 2: MEETING BODIES
# ================================================================

def build_meeting_bodies_tab(app):
    """Build the meeting bodies configuration tab."""
    tab = app.tab_frames["Meeting Bodies"]
    
    # Header with controls
    header = ctk.CTkFrame(tab)
    header.pack(fill="x", padx=10, pady=10)
    
    # Base directory selection
    ctk.CTkLabel(header, text="Base Directory:", 
                font=ctk.CTkFont(size=14, weight="bold")).pack(side="left", padx=5)
    
    base_entry = ctk.CTkEntry(header, textvariable=app.base_dir, width=400,
                             placeholder_text="Parent folder containing /meetings subfolder")
    base_entry.pack(side="left", padx=5)
    
    ctk.CTkButton(header, text="Browse", width=100,
                 command=app._browse_base_dir).pack(side="left", padx=5)
    
    ctk.CTkButton(header, text="Auto-Scan", width=100,
                 command=app._auto_scan_meetings).pack(side="left", padx=5)
    
    # Date format (global)
    ctk.CTkLabel(header, text="Date Format:").pack(side="left", padx=(20, 5))
    ctk.CTkOptionMenu(header, variable=app.date_format, width=120,
                     values=["US DATE", "EU DATE"]).pack(side="left", padx=5)
    
    # Action buttons
    button_frame = ctk.CTkFrame(tab)
    button_frame.pack(fill="x", padx=10, pady=5)
    
    ctk.CTkButton(button_frame, text="+ Add Body", width=120,
                 command=app._add_meeting_body).pack(side="left", padx=5)
    
    ctk.CTkButton(button_frame, text="Remove Selected", width=120,
                 command=app._remove_selected_bodies).pack(side="left", padx=5)
    
    ctk.CTkButton(button_frame, text="Clear All", width=100,
                 fg_color="#c44536", hover_color="#a03529",
                 command=app._clear_all_meeting_bodies).pack(side="left", padx=5)
    
    ctk.CTkButton(button_frame, text="📥 Import from Sheet", width=150,
                 fg_color="#555555", hover_color="#666666",
                 command=app._import_from_sheet).pack(side="left", padx=5)
    
    # Table headers with grid layout for perfect alignment
    headers_frame = ctk.CTkFrame(tab)
    headers_frame.pack(fill="x", padx=10, pady=5)
    
    # Configure columns to expand proportionally with minimum sizes
    headers_frame.grid_columnconfigure(0, minsize=30)   # Checkbox
    headers_frame.grid_columnconfigure(1, weight=2, minsize=300)  # Folder Path
    headers_frame.grid_columnconfigure(2, weight=1, minsize=180)  # Meeting Type
    headers_frame.grid_columnconfigure(3, weight=1, minsize=180)  # Event Category
    headers_frame.grid_columnconfigure(4, minsize=135)  # Default Time (fixed)
    
    ctk.CTkLabel(headers_frame, text="☐", width=30, anchor="w",
                font=ctk.CTkFont(weight="bold")).grid(row=0, column=0, padx=2, sticky="ew")
    ctk.CTkLabel(headers_frame, text="Folder Path", anchor="w",
                font=ctk.CTkFont(weight="bold")).grid(row=0, column=1, padx=2, sticky="ew")
    ctk.CTkLabel(headers_frame, text="Meeting Type", anchor="w",
                font=ctk.CTkFont(weight="bold")).grid(row=0, column=2, padx=2, sticky="ew")
    ctk.CTkLabel(headers_frame, text="Event Category", anchor="w",
                font=ctk.CTkFont(weight="bold")).grid(row=0, column=3, padx=2, sticky="ew")
    ctk.CTkLabel(headers_frame, text="Default Time", width=135, anchor="w",
                font=ctk.CTkFont(weight="bold")).grid(row=0, column=4, padx=2, sticky="ew")
    
    # Scrollable body list
    app.bodies_scroll = ctk.CTkScrollableFrame(tab)
    app.bodies_scroll.pack(fill="both", expand=True, padx=10, pady=10)
    
    # Next button for Meeting Bodies tab
    next_frame = ctk.CTkFrame(tab)
    next_frame.pack(fill="x", side="bottom", padx=10, pady=10)
    
    ctk.CTkButton(
        next_frame,
        text="Next →",
        width=120,
        command=app._next_from_meeting_bodies,
        fg_color="#2b7a4f",
        hover_color="#1e5a3b"
    ).pack(side="right")
    
    app._refresh_meeting_bodies_display()

def browse_base_dir(app):
    """Browse for base directory."""
    folder = filedialog.askdirectory(title="Select Base Directory (Parent Folder)")
    if folder:
        # Save current config before switching
        if app.base_dir.get():
            app._save_config()
        
        # Set new base directory
        app.base_dir.set(folder)
        
        # Try to load config from new directory
        app._load_config()

def import_from_sheet(app):
    """Import Meeting Bodies and Subtypes from Import Sheet.xlsx."""
    import openpyxl
    
    base = app.base_dir.get()
    if not base:
        messagebox.showwarning("No Base Directory", "Please select a base directory first.")
        return
    
    # Look for Import Sheet.xlsx
    sheet_path = Path(base) / "Import Sheet.xlsx"
    if not sheet_path.exists():
        messagebox.showerror("File Not Found", 
            f"Import Sheet.xlsx not found in:\n{base}\n\nPlease make sure the file exists.")
        return
    
    try:
        wb = openpyxl.load_workbook(sheet_path, data_only=True)
        
        # Import Meeting Bodies from "Meeting Bodies" sheet
        if "Meeting Bodies" in wb.sheetnames:
            ws = wb["Meeting Bodies"]
            imported_bodies = []
            
            # Skip header row, read data
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row[0]:  # Skip empty rows
                    continue
                
                folder_path = row[0] or ""
                meeting_type = row[1] or ""
                event_category = row[2] or ""
                default_time = row[3] or ""
                # Note: row[4] is Default Time Text, row[5] is Date Format (not used in MeetingBody)
                
                if folder_path or event_category:  # At least one field populated
                    body = MeetingBody(
                        folder_path=folder_path,
                        meeting_type=meeting_type,
                        event_category=event_category,
                        default_time=str(default_time) if default_time else ""
                    )
                    imported_bodies.append(body)
            
            if imported_bodies:
                app.meeting_bodies = imported_bodies
                app._refresh_meeting_bodies_display()
                messagebox.showinfo("Success", 
                    f"Imported {len(imported_bodies)} Meeting Bodies from Import Sheet!")
        
        # Import Subtypes from "Subtypes" sheet
        if "Subtypes" in wb.sheetnames:
            ws = wb["Subtypes"]
            imported_subtypes = []
            
            # Skip header row, read data
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row[0]:  # Skip empty rows
                    continue
                
                event_category = row[0] or ""
                folder_name = row[1] or ""
                subtype = row[2] or ""
                aliases = row[3] or ""
                override_time = row[4] or ""
                override_meeting_type = row[5] or ""
                event_category_override = row[6] or ""
                
                if event_category or subtype:  # At least one field populated
                    st = Subtype(
                        event_category=event_category,
                        folder_name=folder_name,
                        subtype=subtype,
                        aliases=aliases,
                        override_time=str(override_time) if override_time else "",
                        override_meeting_type=override_meeting_type,
                        event_category_override=event_category_override
                    )
                    imported_subtypes.append(st)
            
            if imported_subtypes:
                app.subtypes = imported_subtypes
                app._refresh_subtypes_display()
                messagebox.showinfo("Success", 
                    f"Imported {len(imported_subtypes)} Subtypes from Import Sheet!")
        
        wb.close()
        
        # Save the imported config
        app._save_config()
        
    except Exception as e:
        messagebox.showerror("Import Error", f"Error importing from sheet:\n\n{str(e)}")

def auto_scan_meetings(app):
    """Auto-scan /meetings folder and populate meeting bodies."""
    base = app.base_dir.get()
    if not base:
        messagebox.showwarning("No Base Directory", "Please select a base directory first.")
        return
    
    # Check for /meetings folder (case-insensitive)
    meetings_dir = None
    base_path = Path(base)
    for item in base_path.iterdir():
        if item.is_dir() and item.name.lower() == "meetings":
            meetings_dir = item
            break
    
    if not meetings_dir:
        messagebox.showerror("Folder Not Found", 
                           f"Could not find /meetings folder in:\n{base}")
        return
    
    # Scan one level deep
    folders = [f for f in meetings_dir.iterdir() if f.is_dir()]
    
    if not folders:
        messagebox.showinfo("No Folders", "No subfolders found in /meetings directory.")
        return
    
    # Add each folder as a meeting body
    for folder in folders:
        # Extract meeting type from folder name
        meeting_type = folder.name.replace("_", " ").replace("-", " ").title()
        
        body = MeetingBody(
            folder_path=str(folder),
            meeting_type=meeting_type,
            event_category=meeting_type,  # Default to same as meeting type
            default_time=""
        )
        app.meeting_bodies.append(body)
    
    app._refresh_meeting_bodies_display()
    messagebox.showinfo("Success", f"Added {len(folders)} meeting bodies from /meetings folder.")

def add_meeting_body(app):
    """Add a new meeting body row."""
    app.meeting_bodies.append(MeetingBody())
    app._refresh_meeting_bodies_display()
    app._save_config()

def remove_selected_bodies(app):
    """Remove selected meeting bodies."""
    to_remove = [i for i, body in enumerate(app.meeting_bodies) if hasattr(body, '_selected') and body._selected]
    for idx in reversed(to_remove):
        app.meeting_bodies.pop(idx)
    app._refresh_meeting_bodies_display()
    app._save_config()

def clear_all_meeting_bodies(app):
    """Clear all meeting bodies after confirmation."""
    if not app.meeting_bodies:
        messagebox.showinfo("Nothing to Clear", "No meeting bodies to clear.")
        return
    
    response = messagebox.askyesno(
        "Clear All Meeting Bodies?",
        f"This will remove all {len(app.meeting_bodies)} meeting bodies.\n\nAre you sure?"
    )
    if response:
        app.meeting_bodies.clear()
        app._refresh_meeting_bodies_display()
        app._save_config()
        messagebox.showinfo("Cleared", "All meeting bodies have been cleared.")

def refresh_meeting_bodies_display(app):
    """Rebuild the meeting bodies list display."""
    # Clear existing
    for widget in app.bodies_scroll.winfo_children():
        widget.destroy()
    
    # Create row for each body
    for idx, body in enumerate(app.meeting_bodies):
        # Alternating row colors
        row_color = "#2b2b2b" if idx % 2 == 0 else "#333333"
        
        row = ctk.CTkFrame(app.bodies_scroll, fg_color=row_color)
        row.pack(fill="x", pady=2)
        
        # Configure columns same as headers
        row.grid_columnconfigure(0, minsize=30)   # Checkbox
        row.grid_columnconfigure(1, weight=2, minsize=300)  # Folder Path
        row.grid_columnconfigure(2, weight=1, minsize=180)  # Meeting Type
        row.grid_columnconfigure(3, weight=1, minsize=180)  # Event Category
        row.grid_columnconfigure(4, minsize=135)  # Default Time (fixed)
        
        # Checkbox for selection
        var = ctk.BooleanVar(value=False)
        checkbox = ctk.CTkCheckBox(row, text="", variable=var, width=30,
                                   command=lambda i=idx, v=var: setattr(app.meeting_bodies[i], '_selected', v.get()))
        checkbox.grid(row=0, column=0, padx=2, sticky="w")
        
        # Folder path with browse
        path_frame = ctk.CTkFrame(row, fg_color="transparent")
        path_frame.grid(row=0, column=1, padx=2, sticky="ew")
        
        path_var = ctk.StringVar(value=body.folder_path)
        path_entry = ctk.CTkEntry(path_frame, textvariable=path_var)
        path_entry.pack(side="left", fill="x", expand=True)
        path_var.trace_add("write", lambda *args, i=idx, v=path_var: 
                         (setattr(app.meeting_bodies[i], 'folder_path', v.get()), app._save_config())[0])
        
        ctk.CTkButton(path_frame, text="...", width=30,
                     command=lambda i=idx: app._browse_body_folder(i)).pack(side="left", padx=(2,0))
        
        # Meeting Type (expandable)
        mt_var = ctk.StringVar(value=body.meeting_type)
        mt_entry = ctk.CTkEntry(row, textvariable=mt_var)
        mt_entry.grid(row=0, column=2, padx=2, sticky="ew")
        mt_var.trace_add("write", lambda *args, i=idx, v=mt_var:
                       (setattr(app.meeting_bodies[i], 'meeting_type', v.get()), app._save_config())[0])
        
        # Event Category (expandable)
        ec_var = ctk.StringVar(value=body.event_category)
        ec_entry = ctk.CTkEntry(row, textvariable=ec_var)
        ec_entry.grid(row=0, column=3, padx=2, sticky="ew")
        ec_var.trace_add("write", lambda *args, i=idx, v=ec_var:
                       (setattr(app.meeting_bodies[i], 'event_category', v.get()), app._save_config())[0])
        
        # Default Time - Time Picker (fixed width)
        time_frame = ctk.CTkFrame(row, fg_color="transparent")
        time_frame.grid(row=0, column=4, padx=2, sticky="w")
        
        # Parse existing time or use default
        current_time = body.default_time if body.default_time else "00:00"
        hour = current_time.split(":")[0] if ":" in current_time else "00"
        minute = current_time.split(":")[1] if ":" in current_time and len(current_time.split(":")) > 1 else "00"
        
        # Hour dropdown (00-23)
        hour_var = ctk.StringVar(value=hour)
        hour_dropdown = ctk.CTkOptionMenu(
            time_frame,
            variable=hour_var,
            values=[f"{h:02d}" for h in range(24)],
            width=60,
            command=lambda val, i=idx: app._update_meeting_time(i)
        )
        hour_dropdown.pack(side="left", padx=1)
        
        ctk.CTkLabel(time_frame, text=":", font=ctk.CTkFont(size=14, weight="bold")).pack(side="left")
        
        # Minute dropdown (00, 15, 30, 45)
        minute_var = ctk.StringVar(value=minute)
        minute_dropdown = ctk.CTkOptionMenu(
            time_frame,
            variable=minute_var,
            values=["00", "15", "30", "45"],
            width=60,
            command=lambda val, i=idx: app._update_meeting_time(i)
        )
        minute_dropdown.pack(side="left", padx=1)
        
        # Store references for updating
        body._hour_var = hour_var
        body._minute_var = minute_var

def update_meeting_time(app, index: int):
    """Update meeting body default time from hour/minute dropdowns."""
    if 0 <= index < len(app.meeting_bodies):
        body = app.meeting_bodies[index]
        if hasattr(body, '_hour_var') and hasattr(body, '_minute_var'):
            hour = body._hour_var.get()
            minute = body._minute_var.get()
            body.default_time = f"{hour}:{minute}"
            app._save_config()

def browse_body_folder(app, index: int):
    """Browse for a specific meeting body folder."""
    folder = filedialog.askdirectory(title="Select Meeting Body Folder")
    if folder and 0 <= index < len(app.meeting_bodies):
        app.meeting_bodies[index].folder_path = folder
        app._refresh_meeting_bodies_display()

